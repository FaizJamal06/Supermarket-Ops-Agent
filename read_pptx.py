from pptx import Presentation
import sys

def read_pptx(filepath):
    prs = Presentation(filepath)
    for i, slide in enumerate(prs.slides):
        print(f"--- Slide {i+1} ---")
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                print(shape.text)

if __name__ == "__main__":
    read_pptx(sys.argv[1])
